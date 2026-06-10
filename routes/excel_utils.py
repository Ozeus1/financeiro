"""Utilitários para geração de planilhas Excel exportadas pelo sistema."""
from io import BytesIO
from datetime import datetime

import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def gerar_excel_relatorio(df, titulo_relatorio, sheet_name, usuario):
    """Gera um arquivo Excel em memória com cabeçalho do sistema, título do
    relatório, usuário e data/hora de geração, seguido pela tabela de dados.
    """
    output = BytesIO()
    linha_inicio_dados = 5  # linha (0-based) onde o cabeçalho da tabela começa

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=linha_inicio_dados)

        ws = writer.sheets[sheet_name]
        num_colunas = max(len(df.columns), 1)
        ultima_coluna = get_column_letter(num_colunas)

        ws['A1'] = 'FiNan - Sistema de Gerenciamento Financeiro'
        ws['A1'].font = Font(size=14, bold=True, color='4361EE')

        ws['A2'] = titulo_relatorio
        ws['A2'].font = Font(size=11, bold=True)

        ws['A3'] = f'Usuário: {usuario}'
        ws['A3'].font = Font(size=9, color='64748B')

        ws['A4'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
        ws['A4'].font = Font(size=9, color='64748B')

        for row in (1, 2, 3, 4):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_colunas)

        # Estiliza o cabeçalho da tabela
        header_fill = PatternFill(start_color='4361EE', end_color='4361EE', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for col_idx in range(1, num_colunas + 1):
            cell = ws.cell(row=linha_inicio_dados + 1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # Ajusta largura das colunas
        for col_idx, coluna in enumerate(df.columns, start=1):
            tamanho = max(len(str(coluna)), df[coluna].astype(str).map(len).max() if not df.empty else 0) + 2
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(tamanho, 12), 40)

    output.seek(0)
    return output
